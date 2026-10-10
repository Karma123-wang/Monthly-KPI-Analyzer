"""Power BI-style PDF, PowerPoint and Excel reports for the monthly inpatient KPIs."""
from __future__ import annotations

import io
import textwrap

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
import pandas as pd
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from kpi_calc import TARGETS, status

# ---------------- Power BI theme ----------------
BLUE, NAVY, RED, GREEN = "#118DFF", "#12239E", "#D64550", "#1AAB40"
ORANGE, GREY, CANVAS, INK, MUTED, GRID = "#E66C37", "#E6E6E6", "#F3F2F1", "#252423", "#605E5C", "#EDEBE9"
SUBTEXT, LABEL = "#E3F1FF", "#CFE6FF"
plt.rcParams.update({
    "font.family": "DejaVu Sans", "font.size": 10, "text.color": INK, "axes.edgecolor": GRID,
    "axes.labelcolor": MUTED, "xtick.color": MUTED, "ytick.color": MUTED,
    "axes.spines.top": False, "axes.spines.right": False, "figure.facecolor": "white",
})
PAGE_W, PAGE_H = landscape(A4)
MARGIN = 28
ALL_SECTIONS = ["summary", "notes", "table", "trends", "wards", "er"]


def _clean(ax):
    ax.spines["left"].set_visible(False)
    ax.spines["bottom"].set_color(GRID)
    ax.tick_params(length=0)
    ax.yaxis.grid(True, color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)


def _kpi_line(ax, months, values, title, target_text, ok, unit, band=None, line=None):
    colours = [GREEN if ok(v) else RED for v in values]
    ax.plot(months, values, color=BLUE, linewidth=2.2, zorder=2)
    ax.scatter(months, values, color=colours, s=46, zorder=3, edgecolor="white", linewidth=1.2)
    for x, v, col in zip(months, values, colours):
        ax.annotate(f"{v:.1f}" if unit == "days" else f"{v:.0f}%", (x, v), textcoords="offset points",
                    xytext=(0, 8), ha="center", fontsize=8.5, color=col, fontweight="bold")
    if band:
        ax.axhspan(band[0], band[1], color=GREEN, alpha=0.10, zorder=0)
    if line is not None:
        ax.axhline(line, color=GREEN, linestyle="--", linewidth=1.3, zorder=1)
    ax.set_title(f"{title}\n", loc="left", fontsize=11.5, fontweight="bold", color=INK)
    ax.text(0, 1.02, f"Target {target_text}  ·  green = met, red = not met", transform=ax.transAxes,
            fontsize=8.5, color=MUTED)
    lo = min(min(values), band[0] if band else line if line is not None else min(values))
    hi = max(max(values), band[1] if band else line if line is not None else max(values))
    pad = (hi - lo) * 0.25 + (1 if unit == "days" else 5)
    ax.set_ylim(max(0, lo - pad), hi + pad)
    if unit == "%":
        ax.yaxis.set_major_formatter(mticker.PercentFormatter(decimals=0))
    ax.tick_params(axis="x", labelsize=8.5, rotation=0)
    _clean(ax)


def _short_month(lbl):
    return lbl.replace(" 20", " '")


def chart_pages(m: pd.DataFrame, wards: pd.DataFrame, hours: pd.Series, sections=None, er=None):
    """Yield (title, subtitle, figure) for each chart page — shared by the PDF and PowerPoint."""
    secs = set(sections or ALL_SECTIONS)
    months = [_short_month(x) for x in m.index]

    # ---- Monthly KPI table ----
    if "table" in secs:
        rows = [("Admissions", "Admissions", "{:,.0f}"), ("Discharges", "Discharges", "{:,.0f}"),
                ("Referred out", "Referred to other facility", "{:,.0f}"), ("Referral %", "Referral % of discharges", "{:.1f}%"),
                ("KPI-56 Discharge 9–11 am %", "KPI-56 Discharged 9–11 am (>90%)", "{:.1f}%"),
                ("KPI-50 ALOS (days)", "KPI-50 ALOS (≤6 days)", "{:.1f}"),
                ("Patient days", "Patient days", "{:,.0f}"),
                ("KPI-51 Bed occupancy %", "KPI-51 Bed occupancy (70–80%)", "{:.1f}%"),
                ("Deaths", "Deaths", "{:,.0f}")]
        cell = [[fmt.format(m.loc[mo, col]) for mo in m.index] for col, _, fmt in rows]
        colours = [[("#FBE3E5" if status(col, m.loc[mo, col]) is False else
                     "#E3F5E8" if status(col, m.loc[mo, col]) else "white") for mo in m.index] for col, _, _ in rows]
        fig, ax = plt.subplots(figsize=(13, 0.5 * len(rows) + 1.1))
        ax.axis("off")
        t = ax.table(cellText=cell, rowLabels=[r[1] for r in rows], colLabels=months, cellColours=colours,
                     loc="center", cellLoc="center", rowLoc="left")
        t.auto_set_font_size(False)
        t.set_fontsize(10)
        t.scale(1, 1.75)
        for (r, c), cl in t.get_celld().items():
            cl.set_edgecolor(GRID)
            if r == 0:
                cl.set_facecolor(BLUE)
                cl.get_text().set_color("white")
                cl.get_text().set_fontweight("bold")
            if c == -1:
                cl.get_text().set_fontweight("bold")
                cl.set_facecolor("#F7F7F7")
            for col, _, _ in rows:
                pass
        for i, (col, _, _) in enumerate(rows, start=1):
            for j, mo in enumerate(m.index):
                if status(col, m.loc[mo, col]) is False:
                    t[i, j].get_text().set_color(RED)
                    t[i, j].get_text().set_fontweight("bold")
        fig.tight_layout()
        yield "Monthly KPI table", "Red = target not met · Green = target met", fig

    # ---- Trend charts for the three BHSQA KPIs ----
    if "trends" in secs:
        fig, axes = plt.subplots(1, 3, figsize=(15, 4.6))
        k56, k50, k51 = (m[k].tolist() for k in TARGETS)
        _kpi_line(axes[0], months, k56, "KPI-56 Discharged 9–11 am", "> 90%", lambda v: v > 90, "%", line=90)
        _kpi_line(axes[1], months, k50, "KPI-50 Average length of stay", "≤ 6 days", lambda v: v <= 6, "days", line=6)
        _kpi_line(axes[2], months, k51, "KPI-51 Bed occupancy", "70–80%", lambda v: 70 <= v <= 80, "%", band=(70, 80))
        fig.tight_layout(w_pad=3)
        yield "KPI trends against BHSQA targets", "Green dashed line / band = target", fig

        fig, axes = plt.subplots(1, 2, figsize=(15, 4.8), gridspec_kw={"width_ratios": [1.35, 1]})
        x = range(len(months))
        w = 0.38
        axes[0].bar([i - w / 2 for i in x], m["Admissions"], w, color=BLUE, label="Admissions")
        axes[0].bar([i + w / 2 for i in x], m["Discharges"], w, color=NAVY, label="Discharges")
        for i, (a, d) in enumerate(zip(m["Admissions"], m["Discharges"])):
            axes[0].text(i - w / 2, a + 1, str(a), ha="center", fontsize=8, color=INK)
            axes[0].text(i + w / 2, d + 1, str(d), ha="center", fontsize=8, color=INK)
        axes[0].set_xticks(list(x), months, fontsize=8.5)
        axes[0].set_title("Admissions and discharges\n", loc="left", fontsize=11.5, fontweight="bold")
        axes[0].legend(loc="lower left", bbox_to_anchor=(0, 1.0), ncol=2, frameon=False, fontsize=9)
        _clean(axes[0])
        ax2 = axes[1]
        ax2.bar(months, m["Referred out"], color=ORANGE, width=0.6)
        for i, (n, p) in enumerate(zip(m["Referred out"], m["Referral %"])):
            ax2.text(i, n + 0.2, f"{n}\n({p:.0f}%)", ha="center", fontsize=8, color=INK)
        ax2.set_ylim(0, max(1, m["Referred out"].max()) * 1.45)
        ax2.set_title("Referred to other facility\n", loc="left", fontsize=11.5, fontweight="bold")
        ax2.text(0, 1.02, "Number of patients (% of discharges)", transform=ax2.transAxes, fontsize=8.5, color=MUTED)
        ax2.tick_params(axis="x", labelsize=8.5)
        _clean(ax2)
        fig.tight_layout(w_pad=3)
        yield "Patient flow and referrals", "Monthly admissions, discharges and referrals out", fig

        fig, ax = plt.subplots(figsize=(15, 4.4))
        ax.bar(months, m["Patient days"], color=BLUE, width=0.6, label="Patient days")
        ax.plot(months, m["Bed days available"], color=MUTED, linestyle="--", linewidth=1.2, label="Bed days available")
        for i, (pd_, bo) in enumerate(zip(m["Patient days"], m["KPI-51 Bed occupancy %"])):
            ax.text(i, pd_ + 15, f"{pd_:,}\n{bo:.0f}% occ.", ha="center", fontsize=8.5,
                    color=RED if not (70 <= bo <= 80) else GREEN)
        ax.set_ylim(0, m["Bed days available"].max() * 1.15)
        ax.set_title("Patient days vs available bed days\n", loc="left", fontsize=11.5, fontweight="bold")
        ax.legend(loc="lower left", bbox_to_anchor=(0, 1.0), ncol=2, frameon=False, fontsize=9)
        _clean(ax)
        fig.tight_layout()
        yield "Patient days", "Occupied bed days each month (KPI-51 numerator)", fig

    # ---- Wards + discharge timing ----
    if "wards" in secs and not wards.empty:
        wv = wards[(wards["Discharges"] > 0) & wards["ALOS (days)"].notna()].head(8)
        fig, axes = plt.subplots(1, 2, figsize=(15, 4.9), gridspec_kw={"width_ratios": [1.1, 1]})
        names = [n[:22] for n in wv.index][::-1]
        axes[0].barh(names, wv["ALOS (days)"][::-1], color=[RED if v > 6 else BLUE for v in wv["ALOS (days)"][::-1]],
                     height=0.6)
        for i, (v, n) in enumerate(zip(wv["ALOS (days)"][::-1], wv["Discharges"][::-1])):
            axes[0].text(v + 0.1, i, f"{v:.1f} d  (n={n})", va="center", fontsize=8.5)
        axes[0].axvline(6, color=GREEN, linestyle="--", linewidth=1.2)
        axes[0].set_xlim(0, max(7, wv["ALOS (days)"].max() * 1.35))
        axes[0].set_title("ALOS by ward\n", loc="left", fontsize=11.5, fontweight="bold")
        axes[0].text(0, 1.02, "Target ≤ 6 days (dashed) · red = above target", transform=axes[0].transAxes,
                     fontsize=8.5, color=MUTED)
        axes[0].xaxis.grid(True, color=GRID)
        axes[0].spines["left"].set_visible(False)
        axes[0].tick_params(length=0)
        hrs = hours.loc[6:22]
        cols = [GREEN if 9 <= h < 11 else (ORANGE if h >= 11 else MUTED) for h in hrs.index]
        axes[1].bar([f"{h}" for h in hrs.index], hrs.values, color=cols, width=0.75)
        axes[1].set_title("Discharges by hour of day\n", loc="left", fontsize=11.5, fontweight="bold")
        total = max(1, hours.sum())
        axes[1].text(0, 1.02, f"Green = 9–11 am ({hours.loc[9:10].sum() / total * 100:.0f}%) · "
                              f"orange = after 11 am ({hours.loc[11:].sum() / total * 100:.0f}%)",
                     transform=axes[1].transAxes, fontsize=8.5, color=MUTED)
        axes[1].set_xlabel("Hour (24-h clock)")
        _clean(axes[1])
        fig.tight_layout(w_pad=3)
        yield "Wards and discharge timing", "For the selected period", fig

    # ---- Emergency observation (reported separately) ----
    if "er" in secs and er is not None and not er.empty:
        e = er.loc[[x for x in m.index if x in er.index]]
        if not e.empty:
            em = [_short_month(x) for x in e.index]
            fig = plt.figure(figsize=(15, 6.2))
            gs = fig.add_gridspec(2, 2, height_ratios=[1.15, 1], width_ratios=[1.35, 1], hspace=0.55, wspace=0.18)
            ax = fig.add_subplot(gs[0, 0])
            x = range(len(em))
            ax.bar(x, e["Admissions"], color="#744EC2", width=0.6)
            for i, v in enumerate(e["Admissions"]):
                ax.text(i, v + 1, str(v), ha="center", fontsize=8.5)
            ax.set_xticks(list(x), em, fontsize=8.5)
            ax.set_ylim(0, e["Admissions"].max() * 1.2)
            ax.set_title("Emergency observation admissions\n", loc="left", fontsize=11.5, fontweight="bold")
            _clean(ax)
            ax2 = fig.add_subplot(gs[0, 1])
            ax2.bar(em, e["Referred out"], color=ORANGE, width=0.6)
            for i, (n, p) in enumerate(zip(e["Referred out"], e["Referral %"])):
                ax2.text(i, n + 0.3, f"{n}\n({p:.0f}%)", ha="center", fontsize=8)
            ax2.set_ylim(0, max(1, e["Referred out"].max()) * 1.5)
            ax2.set_title("Referred to other facility\n", loc="left", fontsize=11.5, fontweight="bold")
            ax2.tick_params(axis="x", labelsize=8.5)
            _clean(ax2)
            axt = fig.add_subplot(gs[1, :])
            axt.axis("off")
            rows = ["Admissions", "Discharges", "Referred out", "Sent to ward", "Deaths", "LAMA / absconded",
                    "Median stay (hours)", "Stayed > 24 hours"]
            cell = [[(f"{e.loc[mo, rw]:.1f}" if rw == "Median stay (hours)" else f"{int(e.loc[mo, rw])}")
                     for mo in e.index] for rw in rows]
            t = axt.table(cellText=cell, rowLabels=rows, colLabels=em, loc="center", cellLoc="center")
            t.auto_set_font_size(False)
            t.set_fontsize(9)
            t.scale(1, 1.35)
            for (rr, cc), cl in t.get_celld().items():
                cl.set_edgecolor(GRID)
                if rr == 0:
                    cl.set_facecolor("#744EC2")
                    cl.get_text().set_color("white")
                    cl.get_text().set_fontweight("bold")
                if cc == -1:
                    cl.get_text().set_fontweight("bold")
                    cl.set_facecolor("#F7F7F7")
            yield ("Emergency observation (ER IP)",
                   "Reported separately – not included in the ward KPIs (KPI-50, 51, 56)", fig)


# ---------------- KPI cards ----------------

def kpi_cards(m: pd.DataFrame, month: str, er=None):
    r = m.loc[month]
    cards = [(f"{int(r['Admissions']):,}", "Admissions", None), (f"{int(r['Discharges']):,}", "Discharges", None),
             (f"{int(r['Referred out'])}", f"Referred to other facility ({r['Referral %']:.1f}%)", None)]
    for kpi, (label, unit, target, _) in TARGETS.items():
        v = r[kpi]
        extra = (f"{int(r['Discharged 9–11 am'])} of {int(r['Discharges'])} patients discharged 9–11 am  ·  "
                 if kpi.startswith("KPI-56") else "")
        cards.append((f"{v:.1f}" + ("%" if unit == "%" else " d"), f"{label}  ·  {extra}target {target}",
                      status(kpi, v)))
    cards.append((f"{int(r['Patient days']):,}", "Patient days", None))
    if er is not None and month in er.index:
        e = int(er.loc[month, "Admissions"])
        cards.append((f"{int(r['Admissions']) + e:,}", f"All admissions (ward {int(r['Admissions'])} + "
                                                        f"emergency observation {e})", None))
    return cards


# ---------------- PDF ----------------

def _wrap(text, width, max_lines):
    lines = textwrap.wrap(str(text), max(10, width))
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] = lines[-1][:-1] + "…"
    return lines


def _bg(c):
    c.setFillColor(CANVAS)
    c.rect(0, 0, PAGE_W, PAGE_H, fill=1, stroke=0)


def _header(c, title, subtitle):
    c.setFillColor(BLUE)
    c.rect(0, PAGE_H - 62, PAGE_W, 62, fill=1, stroke=0)
    c.setFillColor(NAVY)
    c.rect(0, PAGE_H - 62, 6, 62, fill=1, stroke=0)
    c.setFillColor("white")
    c.setFont("Helvetica-Bold", 20)
    c.drawString(MARGIN, PAGE_H - 34, title)
    if subtitle:
        c.setFillColor(SUBTEXT)
        c.setFont("Helvetica", 10.5)
        c.drawString(MARGIN, PAGE_H - 51, subtitle)


def _footer(c, text):
    c.setFillColor(MUTED)
    c.setFont("Helvetica", 8.5)
    c.drawString(MARGIN, 14, text)
    c.drawRightString(PAGE_W - MARGIN, 14, f"Page {c.getPageNumber()}")


def _card(c, x, y, w, h, top=None):
    c.setFillColor("#E1DFDD")
    c.roundRect(x + 1, y - 1.5, w, h, 5, fill=1, stroke=0)
    c.setFillColor("white")
    c.roundRect(x, y, w, h, 5, fill=1, stroke=0)
    if top:
        c.setFillColor(top)
        c.rect(x + 2, y + h - 4, w - 4, 4, fill=1, stroke=0)


def _fig_img(fig):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return ImageReader(buf)


def _cover(c, title, role, author, period_line):
    c.setFillColor(BLUE)
    c.rect(0, 0, PAGE_W, PAGE_H, fill=1, stroke=0)
    c.saveState()
    c.setFillColor("white")
    for cx, cy, r, a in [(PAGE_W - 30, PAGE_H - 20, 240, 0.07), (PAGE_W - 120, 40, 170, 0.06),
                         (PAGE_W - 250, PAGE_H / 2 - 10, 90, 0.05)]:
        c.setFillAlpha(a)
        c.circle(cx, cy, r, fill=1, stroke=0)
    c.restoreState()
    c.setFillColor(NAVY)
    c.rect(0, 0, 16, PAGE_H, fill=1, stroke=0)
    x0 = 64
    c.setFillColor(LABEL)
    c.setFont("Helvetica-Bold", 10.5)
    c.drawString(x0, PAGE_H - 70, "R E P O R T")
    c.setFillColor("white")
    c.rect(x0, PAGE_H - 92, 64, 4, fill=1, stroke=0)
    size = 40 if len(title) <= 36 else 34
    y = PAGE_H - 92 - 22 - size
    c.setFont("Helvetica-Bold", size)
    for row in _wrap(title, 32 if size == 40 else 38, 3):
        c.drawString(x0, y, row)
        y -= size + 8
    c.setFillColor(SUBTEXT)
    c.setFont("Helvetica", 13)
    c.drawString(x0, y - 4, period_line[:120])
    base = 70
    c.setStrokeColor("white")
    c.setStrokeAlpha(0.55)
    c.setLineWidth(0.8)
    c.line(x0, base + 92, x0 + 300, base + 92)
    c.setStrokeAlpha(1)
    if author or role:
        c.setFillColor(LABEL)
        c.setFont("Helvetica-Bold", 9.5)
        c.drawString(x0, base + 70, "P R E P A R E D   B Y")
    if author:
        c.setFillColor("white")
        c.setFont("Helvetica-Bold", 22)
        c.drawString(x0, base + 40, author)
    if role:
        c.setFillColor(SUBTEXT)
        c.setFont("Helvetica", 14)
        c.drawString(x0, base + (16 if author else 40), role)


def build_pdf(m, wards, hours, month, cover, notes=None, sections=None, data_notes=None, er=None) -> bytes:
    """cover = dict(title, role, author, period_line)."""
    secs = set(sections or ALL_SECTIONS)
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=landscape(A4))
    footer = " – ".join(x for x in (cover.get("title"), cover.get("role")) if x)
    _cover(c, cover.get("title", ""), cover.get("role", ""), cover.get("author", ""), cover.get("period_line", ""))
    c.showPage()

    if "summary" in secs:
        _bg(c)
        _header(c, f"Summary – {month}", "Key indicators for the month · red = target not met · green = target met")
        cards = kpi_cards(m, month, er)
        per_row, gap = 4, 12
        kw = (PAGE_W - 2 * MARGIN - gap * (per_row - 1)) / per_row
        kh = 100
        for i, (value, label, ok) in enumerate(cards):
            row, col = divmod(i, per_row)
            x = MARGIN + col * (kw + gap)
            y = PAGE_H - 62 - 18 - kh - row * (kh + gap)
            top = BLUE if ok is None else (GREEN if ok else RED)
            _card(c, x, y, kw, kh, top)
            c.setFillColor(INK if ok is None else (GREEN if ok else RED))
            c.setFont("Helvetica-Bold", 30)
            c.drawString(x + 16, y + kh - 46, value)
            c.setFillColor(MUTED)
            c.setFont("Helvetica", 10)
            lines = [ln for part in label.split("  ·  ") for ln in _wrap(part, int(kw / 5.2), 2)][:3]
            for j, line in enumerate(lines):
                c.drawString(x + 16, y + kh - 64 - j * 11.5, line)
            if ok is not None:
                c.setFont("Helvetica-Bold", 9)
                c.setFillColor(GREEN if ok else RED)
                c.drawRightString(x + kw - 12, y + kh - 22, "MET" if ok else "NOT MET")
        if data_notes:
            y0 = PAGE_H - 62 - 18 - 2 * (kh + gap) - 14
            c.setFillColor(MUTED)
            c.setFont("Helvetica-Oblique", 9)
            for i, n in enumerate(data_notes[:4]):
                c.drawString(MARGIN, y0 - i * 12, f"Note: {n}"[:170])
        _footer(c, footer)
        c.showPage()

    if "notes" in secs and notes and (notes.get("findings") or notes.get("recommendations")):
        _bg(c)
        _header(c, "Key findings & recommendations", cover.get("period_line", ""))
        boxes = [(h, it) for h, it in (("Key findings", notes.get("findings") or []),
                                       ("Recommendations", notes.get("recommendations") or [])) if it]
        gap = 12
        bw = (PAGE_W - 2 * MARGIN - gap * (len(boxes) - 1)) / len(boxes)
        by, bh = 32, PAGE_H - 62 - 16 - 32
        for bi, (heading, items) in enumerate(boxes):
            x = MARGIN + bi * (bw + gap)
            _card(c, x, by, bw, bh, BLUE if bi == 0 else NAVY)
            c.setFillColor(INK)
            c.setFont("Helvetica-Bold", 15)
            c.drawString(x + 18, by + bh - 34, heading)
            yy = by + bh - 62
            chars = int((bw - 60) / 6.4)
            for n, item in enumerate(items, 1):
                lines = textwrap.wrap(item, chars) or [""]
                if yy - 18 * len(lines) < by + 14:
                    break
                c.setFillColor(BLUE if bi == 0 else NAVY)
                c.setFont("Helvetica-Bold", 13)
                c.drawString(x + 18, yy, f"{n}." if bi == 1 else "•")
                c.setFillColor(INK)
                c.setFont("Helvetica", 13)
                for line in lines:
                    c.drawString(x + 38, yy, line)
                    yy -= 18
                yy -= 9
        _footer(c, footer)
        c.showPage()

    for title, subtitle, fig in chart_pages(m, wards, hours, secs, er):
        _bg(c)
        _header(c, title, subtitle)
        img = _fig_img(fig)
        x, y, w, h = MARGIN, 32, PAGE_W - 2 * MARGIN, PAGE_H - 62 - 16 - 32
        _card(c, x, y, w, h)
        iw, ih = img.getSize()
        s = min((w - 28) / iw, (h - 28) / ih)
        c.drawImage(img, x + (w - iw * s) / 2, y + h - 14 - ih * s, iw * s, ih * s)
        _footer(c, footer)
        c.showPage()
    c.save()
    return buf.getvalue()


# ---------------- PowerPoint ----------------

def build_pptx(m, wards, hours, month, cover, notes=None, sections=None, data_notes=None, er=None) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu, Inches, Pt

    def rgb(h):
        h = h.lstrip("#")
        return RGBColor(int(h[:2], 16), int(h[2:4], 16), int(h[4:], 16))

    secs = set(sections or ALL_SECTIONS)
    W, H, HH = Inches(13.333), Inches(7.5), Inches(0.95)
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    footer = " – ".join(x for x in (cover.get("title"), cover.get("role")) if x)

    def rect(sl, x, y, w, h, col, shape=MSO_SHAPE.RECTANGLE):
        s = sl.shapes.add_shape(shape, x, y, w, h)
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            s.adjustments[0] = 0.04
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(col)
        s.line.fill.background()
        s.shadow.inherit = False
        return s

    def text(sl, t, x, y, w, h, size=14, bold=False, col=INK):
        tb = sl.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(str(t).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.size, p.font.bold, p.font.name = Pt(size), bold, "Segoe UI"
            p.font.color.rgb = rgb(col)
            p.space_after = Pt(6)

    def card(sl, x, y, w, h, top=None):
        rect(sl, x + Emu(12000), y + Emu(15000), w, h, "#E1DFDD", MSO_SHAPE.ROUNDED_RECTANGLE)
        rect(sl, x, y, w, h, "#FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE)
        if top:
            rect(sl, x + Inches(0.05), y, w - Inches(0.1), Inches(0.06), top)

    def base(title, subtitle):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = rgb(CANVAS)
        rect(sl, 0, 0, W, HH, BLUE)
        rect(sl, 0, 0, Inches(0.09), HH, NAVY)
        text(sl, title, Inches(0.4), Inches(0.14), Inches(12), Inches(0.5), 24, True, "#FFFFFF")
        if subtitle:
            text(sl, subtitle, Inches(0.4), Inches(0.58), Inches(12), Inches(0.3), 12, False, SUBTEXT)
        text(sl, footer, Inches(0.4), Inches(7.15), Inches(11), Inches(0.25), 9, False, MUTED)
        return sl

    # Cover
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.background.fill.solid()
    sl.background.fill.fore_color.rgb = rgb(BLUE)
    for cx, cy, r, col in [(13.0, 0.2, 3.3, "#2A99FF"), (11.9, 7.3, 2.4, "#2696FF"), (9.7, 3.6, 1.25, "#2A99FF")]:
        rect(sl, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r), col, MSO_SHAPE.OVAL)
    rect(sl, 0, 0, Inches(0.22), H, NAVY)
    text(sl, "R E P O R T", Inches(0.9), Inches(0.75), Inches(4), Inches(0.3), 11, True, LABEL)
    rect(sl, Inches(0.9), Inches(1.15), Inches(0.9), Inches(0.06), "#FFFFFF")
    text(sl, cover.get("title", ""), Inches(0.9), Inches(1.4), Inches(8.4), Inches(1.9), 44, True, "#FFFFFF")
    text(sl, cover.get("period_line", ""), Inches(0.9), Inches(3.35), Inches(9), Inches(0.4), 14, False, SUBTEXT)
    rect(sl, Inches(0.9), Inches(5.35), Inches(4.2), Emu(9525), "#FFFFFF")
    if cover.get("author") or cover.get("role"):
        text(sl, "P R E P A R E D   B Y", Inches(0.9), Inches(5.55), Inches(5), Inches(0.3), 10, True, LABEL)
    if cover.get("author"):
        text(sl, cover["author"], Inches(0.9), Inches(5.85), Inches(8), Inches(0.5), 26, True, "#FFFFFF")
    if cover.get("role"):
        text(sl, cover["role"], Inches(0.9), Inches(6.45 if cover.get("author") else 5.85), Inches(8), Inches(0.4),
             16, False, SUBTEXT)

    if "summary" in secs:
        sl = base(f"Summary – {month}", "Key indicators for the month · red = target not met · green = target met")
        cards = kpi_cards(m, month, er)
        left, gap = Inches(0.35), Inches(0.2)
        kw = int((W - 2 * left - gap * 3) / 4)
        kh = Inches(1.65)
        for i, (value, label, ok) in enumerate(cards):
            row, col = divmod(i, 4)
            x = left + col * (kw + gap)
            y = HH + Inches(0.25) + row * (kh + gap)
            card(sl, x, y, kw, kh, BLUE if ok is None else (GREEN if ok else RED))
            text(sl, value, x + Inches(0.2), y + Inches(0.2), kw - Inches(0.3), Inches(0.6), 30, True,
                 INK if ok is None else (GREEN if ok else RED))
            text(sl, label, x + Inches(0.2), y + Inches(0.9), kw - Inches(0.35), Inches(0.6), 11, False, MUTED)
            if ok is not None:
                text(sl, "MET" if ok else "NOT MET", x + kw - Inches(1.2), y + Inches(0.18), Inches(1.0), Inches(0.3),
                     10, True, GREEN if ok else RED)
        if data_notes:
            text(sl, "\n".join("Note: " + n for n in data_notes[:4]), Inches(0.4), Inches(5.5), Inches(12.4),
                 Inches(1.4), 10, False, MUTED)

    if "notes" in secs and notes and (notes.get("findings") or notes.get("recommendations")):
        sl = base("Key findings & recommendations", cover.get("period_line", ""))
        boxes = [(h, it) for h, it in (("Key findings", notes.get("findings") or []),
                                       ("Recommendations", notes.get("recommendations") or [])) if it]
        left, gap = Inches(0.35), Inches(0.2)
        bw = int((W - 2 * left - gap * (len(boxes) - 1)) / len(boxes))
        by = HH + Inches(0.2)
        bh = H - by - Inches(0.45)
        for bi, (heading, items) in enumerate(boxes):
            x = left + bi * (bw + gap)
            card(sl, x, by, bw, bh, BLUE if bi == 0 else NAVY)
            text(sl, heading, x + Inches(0.3), by + Inches(0.25), bw - Inches(0.6), Inches(0.4), 18, True)
            body = "\n".join((f"{n}.  " if bi == 1 else "•  ") + it for n, it in enumerate(items, 1))
            text(sl, body, x + Inches(0.3), by + Inches(0.8), bw - Inches(0.6), bh - Inches(1.0),
                 14 if sum(map(len, items)) < 700 else 12)

    for title, subtitle, fig in chart_pages(m, wards, hours, secs, er):
        sl = base(title, subtitle)
        x, y = Inches(0.35), HH + Inches(0.2)
        w, h = W - Inches(0.7), H - y - Inches(0.45)
        card(sl, x, y, w, h)
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=200, bbox_inches="tight", facecolor="white")
        fw, fh = fig.get_size_inches()
        plt.close(fig)
        buf.seek(0)
        pad = Inches(0.2)
        s = min((w - 2 * pad) / fw, (h - 2 * pad) / fh)
        pw, ph = Emu(int(fw * s)), Emu(int(fh * s))
        sl.shapes.add_picture(buf, int(x + (w - pw) / 2), int(y + pad), pw, ph)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# ---------------- Excel (for submission / further work) ----------------

def build_xlsx(m, wards, cover, data_notes=None, er=None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="D0D7E2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill, head_font = PatternFill("solid", fgColor=BLUE[1:]), Font(bold=True, color="FFFFFF")
    red_fill, green_fill = PatternFill("solid", fgColor="FBE3E5"), PatternFill("solid", fgColor="E3F5E8")
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly KPIs"
    ws["A1"] = cover.get("title", "")
    ws["A1"].font = Font(bold=True, size=15, color=BLUE[1:])
    ws["A2"] = " · ".join(x for x in (cover.get("role"), cover.get("author"), cover.get("period_line")) if x)
    ws["A2"].font = Font(italic=True, color="605E5C")
    cols = ["Admissions", "Discharges", "Referred out", "Referral %", "Discharged 9–11 am", "KPI-56 Discharge 9–11 am %",
            "Inpatient days of discharged", "KPI-50 ALOS (days)", "Patient days", "Bed days available",
            "KPI-51 Bed occupancy %", "Deaths", "LAMA / absconded", "Long stays (> 30 days)"]
    header = ["Month"] + cols
    targets = {"KPI-56 Discharge 9–11 am %": "> 90%", "KPI-50 ALOS (days)": "≤ 6", "KPI-51 Bed occupancy %": "70–80%"}
    for j, h in enumerate(header, 1):
        cl = ws.cell(row=4, column=j, value=h)
        cl.fill, cl.font, cl.border = head_fill, head_font, border
        cl.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        t = ws.cell(row=5, column=j, value=("Target" if j == 1 else targets.get(h, "")))
        t.font, t.border = Font(italic=True, color="605E5C"), border
    for i, mo in enumerate(m.index, start=6):
        ws.cell(row=i, column=1, value=mo).border = border
        for j, col in enumerate(cols, start=2):
            v = m.loc[mo, col]
            cl = ws.cell(row=i, column=j, value=round(float(v), 1) if isinstance(v, float) else int(v))
            cl.border = border
            st = status(col, v)
            if st is not None:
                cl.fill = green_fill if st else red_fill
                cl.font = Font(bold=not st, color="1A1A1A" if st else RED[1:])
    ws.column_dimensions["A"].width = 12
    for j in range(2, len(header) + 1):
        ws.column_dimensions[ws.cell(row=4, column=j).column_letter].width = 14
    ws.row_dimensions[4].height = 46
    ws.freeze_panes = "B6"
    r = 6 + len(m) + 1
    ws.cell(row=r, column=1, value="Definitions (BHSQA KPI user guide)").font = Font(bold=True)
    for k, line in enumerate([
        "KPI-56: patients discharged between 9 am and 11 am ÷ all patients discharged in the month × 100 (target > 90%).",
        "KPI-50: inpatient days of patients discharged in the month ÷ patients discharged (target ≤ 6 days). "
        "Same-day stay = 1 day.",
        "KPI-51: inpatient (patient) days in the month ÷ available bed days × 100 (target 70–80%).",
        "Ward inpatients (IPD) only; emergency observation excluded. Based on Actual Discharge Date & Time.",
    ] + [f"Note: {n}" for n in (data_notes or [])], start=1):
        ws.cell(row=r + k, column=1, value=line)
    w2 = wb.create_sheet("By ward")
    for j, h in enumerate(["Ward"] + list(wards.columns), 1):
        cl = w2.cell(row=1, column=j, value=h)
        cl.fill, cl.font, cl.border = head_fill, head_font, border
    for i, (ward, row) in enumerate(wards.iterrows(), start=2):
        w2.cell(row=i, column=1, value=ward).border = border
        for j, v in enumerate(row, start=2):
            cl = w2.cell(row=i, column=j, value=None if pd.isna(v) else round(float(v), 1))
            cl.border = border
    w2.column_dimensions["A"].width = 24
    if er is not None and not er.empty:
        e = er.loc[[x for x in m.index if x in er.index]]
        w3 = wb.create_sheet("Emergency observation")
        w3["A1"] = "Emergency observation (ER IP) – reported separately, not part of the ward KPIs"
        w3["A1"].font = Font(bold=True, size=12, color=BLUE[1:])
        hdr = ["Month"] + list(e.columns) + ["All admissions (ward + ER obs)"]
        for j, h in enumerate(hdr, 1):
            cl = w3.cell(row=3, column=j, value=h)
            cl.fill, cl.font, cl.border = head_fill, head_font, border
            cl.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        for i, mo in enumerate(e.index, start=4):
            w3.cell(row=i, column=1, value=mo).border = border
            for j, col in enumerate(e.columns, start=2):
                v = e.loc[mo, col]
                w3.cell(row=i, column=j, value=round(float(v), 1) if isinstance(v, float) else int(v)).border = border
            w3.cell(row=i, column=len(hdr), value=int(e.loc[mo, "Admissions"]) + int(m.loc[mo, "Admissions"])).border = border
        w3.row_dimensions[3].height = 46
        w3.column_dimensions["A"].width = 12
        for j in range(2, len(hdr) + 1):
            w3.column_dimensions[w3.cell(row=3, column=j).column_letter].width = 14
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
